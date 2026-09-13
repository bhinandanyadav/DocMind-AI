from dataclasses import dataclass, field
from pathlib import Path
from typing import Optional
import re


@dataclass
class ParsedPage:
    page_number: int
    text: str
    section: Optional[str] = None


@dataclass
class ParsedDocument:
    document_id: str
    document_name: str
    pages: list[ParsedPage] = field(default_factory=list)

    @property
    def total_pages(self) -> int:
        return len(self.pages)

    @property
    def full_text(self) -> str:
        return "\n".join(p.text for p in self.pages)


def _detect_section(text: str) -> Optional[str]:
    """Detect section heading from first non-empty line."""
    lines = [l.strip() for l in text.split("\n") if l.strip()]
    if not lines:
        return None
    first = lines[0]
    # All caps or Title Case short line = likely a heading
    if len(first) < 100 and (first.isupper() or re.match(r'^[A-Z][^.!?]*$', first)):
        return first
    return None


def parse_pdf(file_path: str, document_id: str) -> ParsedDocument:
    """Extract text from PDF using PyMuPDF, preserving page numbers."""
    import fitz  # PyMuPDF

    doc_name = Path(file_path).name
    parsed = ParsedDocument(document_id=document_id, document_name=doc_name)

    pdf = fitz.open(file_path)
    native_pages = [page.get_text("text").strip() for page in pdf]

    # OCR is expensive. Only use it when the native extraction indicates that
    # this is a scanned document, rather than for every sparse text page.
    native_text_length = sum(len(text) for text in native_pages)
    extracted_pages = native_pages
    if native_text_length < 100:
        try:
            import io
            import pytesseract
            from PIL import Image

            extracted_pages = []
            for page, native_text in zip(pdf, native_pages):
                if native_text:
                    extracted_pages.append(native_text)
                    continue
                pix = page.get_pixmap(dpi=150)
                image = Image.open(io.BytesIO(pix.tobytes("png")))
                extracted_pages.append(pytesseract.image_to_string(image).strip())
        except Exception:
            extracted_pages = native_pages

    for page_num, text in enumerate(extracted_pages, start=1):
        if text:
            parsed.pages.append(ParsedPage(
                page_number=page_num,
                text=text,
                section=_detect_section(text),
            ))

    pdf.close()
    return parsed


def parse_docx(file_path: str, document_id: str) -> ParsedDocument:
    """Extract text from DOCX using python-docx, preserving headings as sections."""
    from docx import Document as DocxDocument

    doc_name = Path(file_path).name
    parsed = ParsedDocument(document_id=document_id, document_name=doc_name)

    docx = DocxDocument(file_path)
    current_section = None
    current_text_parts: list[str] = []
    virtual_page = 1
    chars_per_page = 3000
    char_count = 0

    for para in docx.paragraphs:
        text = para.text.strip()
        if not text:
            continue

        # Headings become section markers
        if para.style.name.startswith("Heading"):
            # Save current chunk as a page
            if current_text_parts:
                parsed.pages.append(ParsedPage(
                    page_number=virtual_page,
                    text="\n".join(current_text_parts),
                    section=current_section,
                ))
                virtual_page += 1
                current_text_parts = []
                char_count = 0
            current_section = text
        else:
            current_text_parts.append(text)
            char_count += len(text)
            # Split into virtual pages every ~3000 chars
            if char_count >= chars_per_page:
                parsed.pages.append(ParsedPage(
                    page_number=virtual_page,
                    text="\n".join(current_text_parts),
                    section=current_section,
                ))
                virtual_page += 1
                current_text_parts = []
                char_count = 0

    # Add remaining text
    if current_text_parts:
        parsed.pages.append(ParsedPage(
            page_number=virtual_page,
            text="\n".join(current_text_parts),
            section=current_section,
        ))

    # Fallback: if no pages, add whole doc as one page
    if not parsed.pages:
        full_text = "\n".join(p.text for p in docx.paragraphs if p.text.strip())
        if full_text:
            parsed.pages.append(ParsedPage(page_number=1, text=full_text))

    return parsed


def parse_txt(file_path: str, document_id: str) -> ParsedDocument:
    """Parse plain text file, splitting into virtual pages."""
    doc_name = Path(file_path).name
    parsed = ParsedDocument(document_id=document_id, document_name=doc_name)

    with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
        content = f.read().strip()

    if not content:
        return parsed

    chars_per_page = 3000
    chunks = [content[i:i+chars_per_page] for i in range(0, len(content), chars_per_page)]

    for i, chunk in enumerate(chunks):
        parsed.pages.append(ParsedPage(
            page_number=i + 1,
            text=chunk.strip(),
            section=_detect_section(chunk),
        ))

    return parsed


def parse_document(file_path: str, file_type: str, document_id: str) -> ParsedDocument:
    """Route to correct parser based on file type."""
    file_type = file_type.lower().lstrip(".")
    if file_type == "pdf":
        return parse_pdf(file_path, document_id)
    elif file_type == "docx":
        return parse_docx(file_path, document_id)
    elif file_type == "txt":
        return parse_txt(file_path, document_id)
    else:
        raise ValueError(f"Unsupported file type: {file_type}")
